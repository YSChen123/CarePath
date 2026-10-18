from pathlib import Path
import sys, json, zipfile, xml.etree.ElementTree as ET
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / '.tools/python'))
from pypdf import PdfReader
import pymupdf
from openpyxl import load_workbook

root = Path(__file__).resolve().parents[2]
out = root / 'website/docs/materials'
out.mkdir(parents=True, exist_ok=True)
pdfs = list(root.glob('*.pdf'))
prior = root.parent / '人机交互'
pdfs += [p for p in prior.glob('*.pdf') if '张涛' in p.name and ('-PS.pdf' in p.name or p.name.endswith('202301130286.pdf'))]
for pdf in pdfs:
    reader = PdfReader(pdf)
    text = '\n\n'.join(f'PAGE {i+1}\n{p.extract_text()}' for i,p in enumerate(reader.pages))
    (out / (pdf.stem + '.txt')).write_text(text, encoding='utf-8')
    print('\nPDF',pdf.name, len(reader.pages),'pages\n',text)
    doc = pymupdf.open(pdf)
    for i in range(len(doc)):
        doc[i].get_pixmap(matrix=pymupdf.Matrix(1.3,1.3)).save(out / f'{pdf.stem}-{i+1}.png')
with zipfile.ZipFile(next(root.glob('*.pptx'))) as z:
    slides=[]
    for i in range(1,89):
        r=ET.fromstring(z.read(f'ppt/slides/slide{i}.xml'))
        slides.append(f'SLIDE {i}\n'+'\n'.join(t.text or '' for t in r.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}t')))
    (out/'slides.txt').write_text('\n\n'.join(slides),encoding='utf-8')
wb=load_workbook(next(p for p in root.glob('*.xlsx') if not p.name.startswith('~$')),read_only=True,data_only=True)
for sheet in wb:
    rows=list(sheet.values)
    print('\nEXCEL',sheet.title,sheet.max_row,sheet.max_column)
    (out/(sheet.title+'.json')).write_text(json.dumps(rows,ensure_ascii=False,default=str,indent=2),encoding='utf-8')
seed=json.loads((root/'data/website_seed.json').read_text(encoding='utf-8'))
print('\nINTERVIEWS\n')
for person in seed['interviews']:
    print(person['id'], person['role'], person['scenario'], person['emphasis'])
    for turn in person['turns']:
        print(turn['section'],turn['question_id'],turn['kind'],turn['text'])
