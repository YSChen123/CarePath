import importlib.util, io, unittest
from contextlib import redirect_stdout
from pathlib import Path
from unittest.mock import patch
from openpyxl import load_workbook

site=Path(__file__).resolve().parents[1]
spec=importlib.util.spec_from_file_location('importer',site/'scripts/import-data.py')
importer=importlib.util.module_from_spec(spec)
spec.loader.exec_module(importer)

class ImportValidationTests(unittest.TestCase):
    def check_rejected(self,question,value):
        book=next(p for p in site.parent.glob('*.xlsx') if not p.name.startswith('~$'))
        wb=load_workbook(book,data_only=True)
        sheet=next(s for s in wb if any('Q1' in row and 'Q22' in row for row in s.values))
        header=next(row for row in sheet.values if 'Q1' in row and 'Q22' in row)
        row=next(c.row for c in sheet['A'] if c.value=='R01')
        sheet.cell(row,header.index(question)+1,value)
        with patch.object(importer,'load_workbook',return_value=wb),patch.object(Path,'write_text') as write,redirect_stdout(io.StringIO()):
            with self.assertRaises((ValueError,RuntimeError)):
                importer.main()
            write.assert_not_called()
        wb.close()
    def test_fractional_scale_rejected_before_write(self):
        self.check_rejected('Q5',2.5)
    def test_unknown_option_rejected_before_write(self):
        self.check_rejected('Q1','不是问卷的选项')

if __name__=='__main__': unittest.main()
